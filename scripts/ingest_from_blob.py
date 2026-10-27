"""
Reads every supported document from an Azure Storage Emulator (Azurite)
container, generates a SAS URI for each blob, downloads it, extracts
text based on file extension, and POSTs the extracted text to the
document-addition API (/api/documents), which handles chunking.

Prereqs:
    - Azurite running locally (e.g. via the VS Code extension, or
      `azurite --silent --location ./azurite-data`)
    - A container created and populated with your test documents
    - Your Functions app running separately via `func start`
      (this script talks to it over HTTP, it does not touch Qdrant directly)

Install:
    pip install azure-storage-blob python-docx pypdf openpyxl xlrd python-pptx requests

Usage:
    python scripts/ingest_from_blob.py --container document-container --category Research
"""
import argparse
import io
import logging
import sys
from datetime import datetime, timedelta, timezone

import requests
from azure.storage.blob import BlobServiceClient, generate_blob_sas, BlobSasPermissions

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
log = logging.getLogger("ingest_from_blob")

# Azurite's well-known local development connection string.
# This is a publicly documented fixed key for the local emulator only -
# it is NOT a real credential and is safe to hardcode. Never reuse this
# pattern for a real Azure Storage account's connection string.
AZURITE_CONNECTION_STRING = (
    "DefaultEndpointsProtocol=http;"
    "AccountName=devstoreaccount1;"
    "AccountKey=Eby8vdM02xNOcqFlqUwJPLlmEtlCDXJ1OUzFT50uSRZ6IFsuFq2UVErCz4I6tq/K1SZFPTOtr/KBHBeksoGMGw==;"
    "BlobEndpoint=http://127.0.0.1:10000/devstoreaccount1;"
)

API_BASE_URL = "http://localhost:7071/api"
SAS_EXPIRY_MINUTES = 15

SUPPORTED_EXTENSIONS = {".docx", ".pdf", ".xlsx", ".xls", ".pptx", ".txt", ".md"}


def generate_sas_uri(blob_service_client: BlobServiceClient, container_name: str, blob_name: str) -> str:
    account_name = blob_service_client.account_name
    account_key = blob_service_client.credential.account_key

    sas_token = generate_blob_sas(
        account_name=account_name,
        container_name=container_name,
        blob_name=blob_name,
        account_key=account_key,
        permission=BlobSasPermissions(read=True),
        expiry=datetime.now(timezone.utc) + timedelta(minutes=SAS_EXPIRY_MINUTES),
    )

    blob_client = blob_service_client.get_blob_client(container=container_name, blob=blob_name)
    return f"{blob_client.url}?{sas_token}"


def download_blob_bytes(sas_uri: str) -> bytes:
    resp = requests.get(sas_uri, timeout=60)
    resp.raise_for_status()
    return resp.content


def extract_text_docx(data: bytes) -> str:
    import docx
    document = docx.Document(io.BytesIO(data))
    return "\n".join(p.text for p in document.paragraphs if p.text.strip())


def extract_text_pdf(data: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    pages_text = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            pages_text.append(text)
    return "\n".join(pages_text)


def extract_text_xlsx(data: bytes) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    rows_text = []
    for sheet in wb.worksheets:
        rows_text.append(f"--- Sheet: {sheet.title} ---")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                rows_text.append(" | ".join(cells))
    return "\n".join(rows_text)


def extract_text_xls(data: bytes) -> str:
    # Legacy .xls (pre-2007 binary format) needs xlrd, not openpyxl.
    import xlrd
    wb = xlrd.open_workbook(file_contents=data)
    rows_text = []
    for sheet in wb.sheets():
        rows_text.append(f"--- Sheet: {sheet.name} ---")
        for row_idx in range(sheet.nrows):
            cells = [str(c) for c in sheet.row_values(row_idx) if c != ""]
            if cells:
                rows_text.append(" | ".join(cells))
    return "\n".join(rows_text)


def extract_text_pptx(data: bytes) -> str:
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    slides_text = []
    for i, slide in enumerate(prs.slides, start=1):
        slide_lines = [f"--- Slide {i} ---"]
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_lines.append(shape.text)
        if len(slide_lines) > 1:
            slides_text.append("\n".join(slide_lines))
    return "\n".join(slides_text)


def extract_text_plain(data: bytes) -> str:
    return data.decode("utf-8", errors="replace")


EXTRACTORS = {
    ".docx": extract_text_docx,
    ".pdf": extract_text_pdf,
    ".xlsx": extract_text_xlsx,
    ".xls": extract_text_xls,
    ".pptx": extract_text_pptx,
    ".txt": extract_text_plain,
    ".md": extract_text_plain,
}


def get_extension(blob_name: str) -> str:
    idx = blob_name.rfind(".")
    return blob_name[idx:].lower() if idx != -1 else ""


def extract_text(blob_name: str, data: bytes) -> str:
    ext = get_extension(blob_name)
    extractor = EXTRACTORS.get(ext)
    if extractor is None:
        raise ValueError(f"Unsupported file extension '{ext}' for blob '{blob_name}'")
    return extractor(data)


def post_to_add_document_api(page_content: str, blob_name: str, category: str) -> dict:
    payload = {
        "page_content": page_content,
        "metadata": {
            "category": category,
            "source_filename": blob_name,
        }
    }
    resp = requests.post(
        f"{API_BASE_URL}/documents",
        json=payload,
        timeout=120,
    )
    resp.raise_for_status()
    return resp.json()


def ingest_container(container_name: str, category: str, connection_string: str):
    blob_service_client = BlobServiceClient.from_connection_string(connection_string)
    container_client = blob_service_client.get_container_client(container_name)

    blob_list = list(container_client.list_blobs())
    log.info(f"Found {len(blob_list)} blob(s) in container '{container_name}'.")

    results = {"succeeded": [], "skipped": [], "failed": []}

    for blob in blob_list:
        blob_name = blob.name
        ext = get_extension(blob_name)

        if ext not in SUPPORTED_EXTENSIONS:
            log.warning(f"Skipping '{blob_name}' - unsupported extension '{ext}'.")
            results["skipped"].append(blob_name)
            continue

        try:
            log.info(f"Processing '{blob_name}'...")

            sas_uri = generate_sas_uri(blob_service_client, container_name, blob_name)
            raw_bytes = download_blob_bytes(sas_uri)

            text = extract_text(blob_name, raw_bytes)
            if not text.strip():
                log.warning(f"'{blob_name}' produced no extractable text - skipping.")
                results["skipped"].append(blob_name)
                continue

            api_response = post_to_add_document_api(text, blob_name, category)
            chunks_added = api_response.get("total_chunks_added", "?")
            log.info(f"'{blob_name}' -> added as {chunks_added} chunk(s).")
            results["succeeded"].append(blob_name)

        except Exception as e:
            log.error(f"Failed to process '{blob_name}': {e}")
            results["failed"].append({"blob": blob_name, "error": str(e)})

    log.info("--- Ingestion complete ---")
    log.info(f"Succeeded: {len(results['succeeded'])}")
    log.info(f"Skipped:   {len(results['skipped'])}")
    log.info(f"Failed:    {len(results['failed'])}")
    if results["failed"]:
        for f in results["failed"]:
            log.info(f"  - {f['blob']}: {f['error']}")

    return results


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Ingest documents from an Azurite blob container into the vector DB.")
    parser.add_argument("--container", required=True, help="Name of the Azurite container to read from.")
    parser.add_argument("--category", default="General", help="Metadata category tag to attach to every ingested document.")
    parser.add_argument("--connection-string", default=AZURITE_CONNECTION_STRING, help="Override the default Azurite connection string (e.g. to point at a real Azure Storage account instead).")
    args = parser.parse_args()

    try:
        ingest_container(args.container, args.category, args.connection_string)
    except Exception as e:
        log.error(f"Ingestion run failed: {e}")
        sys.exit(1)